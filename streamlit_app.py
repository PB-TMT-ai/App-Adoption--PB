"""
App Adoption Dashboard — Streamlit Web Interface
Upload your Excel file and view adoption metrics instantly.
"""

import io
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Page Config ──
st.set_page_config(
    page_title="App Adoption Dashboard",
    page_icon="📊",
    layout="wide",
)

# ── Constants ──
VALID_STATUSES = ["Installed", "Not Installed", "Uninstalled", "Tech Issue", "Not Working"]

REQUIRED_COLUMNS = ["Account Sf ID", "Name of the Dealer", "Status"]

STATUS_COLORS = {
    "Installed": "#28a745",
    "Not Installed": "#6c757d",
    "Uninstalled": "#ffc107",
    "Tech Issue": "#fd7e14",
    "Not Working": "#dc3545",
}


# ── Data Loading & Cleaning ──
@st.cache_data
def load_and_clean(uploaded_bytes):
    """Load Excel bytes into a cleaned DataFrame."""
    df = pd.read_excel(io.BytesIO(uploaded_bytes), engine="openpyxl", dtype=str)
    df.columns = [c.strip() for c in df.columns]

    # Clean numeric columns
    int_cols = ["# orders (total)", "# orders (via. app.)"]
    for col in int_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int)

    numeric_cols = ["Order quantity (total)", "Quantity (via. app.)"]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(
                df[col].astype(str).str.strip().replace("-", "0"), errors="coerce"
            ).fillna(0)

    float_cols = ["% orders via. app.", "% order quantity (via. app.)"]
    for col in float_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    return df


def validate_data(df):
    """Validate the DataFrame. Returns (is_valid, errors)."""
    errors = []
    for col in REQUIRED_COLUMNS:
        if col not in df.columns:
            errors.append(f"Missing required column: **{col}**")

    if errors:
        return False, errors

    if "Account Sf ID" in df.columns:
        blank_ids = df["Account Sf ID"].isna().sum()
        if blank_ids > 0:
            errors.append(f"**Account Sf ID** has {blank_ids} blank value(s)")

    if "Status" in df.columns:
        blank_status = df["Status"].isna().sum()
        if blank_status > 0:
            errors.append(f"**Status** has {blank_status} blank value(s)")

    return len(errors) == 0, errors


# ── Report Generation ──
def generate_excel_report(df):
    """Generate the Excel report and return bytes."""
    total = len(df)
    status_counts = df["Status"].value_counts()
    today = datetime.now().strftime("%Y-%m-%d")

    wb = Workbook()

    def style_header(ws):
        hf = Font(bold=True, color="FFFFFF", size=11)
        hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        for cell in ws[1]:
            cell.font = hf
            cell.fill = hfill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.freeze_panes = "A2"

    def auto_width(ws):
        for ci, cc in enumerate(ws.columns, 1):
            ml = max((len(str(c.value or "")) for c in cc), default=8)
            ws.column_dimensions[get_column_letter(ci)].width = min(ml + 3, 45)

    # ── Sheet 1: Summary ──
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    bold = Font(bold=True, size=11)

    installed = int(status_counts.get("Installed", 0))
    not_installed = int(status_counts.get("Not Installed", 0))
    uninstalled = int(status_counts.get("Uninstalled", 0))
    tech_issue = int(status_counts.get("Tech Issue", 0))
    not_working = int(status_counts.get("Not Working", 0))
    adoption_rate = round(installed / total * 100, 1) if total > 0 else 0
    ever_adopted = installed + uninstalled + tech_issue + not_working
    ever_adopted_rate = round(ever_adopted / total * 100, 1) if total > 0 else 0

    metrics = [
        ("Report Date", today),
        ("", ""),
        ("Total Dealers", total),
        ("Installed", installed),
        ("Not Installed", not_installed),
        ("Uninstalled", uninstalled),
        ("Tech Issue", tech_issue),
        ("Not Working", not_working),
        ("", ""),
        ("Currently Installed (%)", f"{adoption_rate}%"),
        ("Ever Adopted (%)", f"{ever_adopted_rate}%"),
    ]
    for label, value in metrics:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold

    style_header(ws)
    auto_width(ws)

    # ── Sheet 2: Zone-wise ──
    ws2 = wb.create_sheet("Zone-wise Adoption")
    ws2.append(["Zone", "State", "Total", "Installed", "Not Installed", "Uninstalled",
                 "Tech Issue", "Not Working", "Adoption Rate (%)"])
    if "Zone" in df.columns and "State" in df.columns:
        zs = df.groupby(["Zone", "State"]).agg(
            total=("Account Sf ID", "count"),
            inst=("Status", lambda x: (x == "Installed").sum()),
            ni=("Status", lambda x: (x == "Not Installed").sum()),
            un=("Status", lambda x: (x == "Uninstalled").sum()),
            ti=("Status", lambda x: (x == "Tech Issue").sum()),
            nw=("Status", lambda x: (x == "Not Working").sum()),
        ).reset_index()
        zs["rate"] = (zs["inst"] / zs["total"] * 100).round(1)
        for _, r in zs.sort_values(["Zone", "rate"], ascending=[True, False]).iterrows():
            ws2.append([r["Zone"], r["State"], int(r["total"]), int(r["inst"]),
                        int(r["ni"]), int(r["un"]), int(r["ti"]), int(r["nw"]), f"{r['rate']}%"])
    style_header(ws2)
    auto_width(ws2)

    # ── Sheet 3: Distributor-wise ──
    ws3 = wb.create_sheet("Distributor-wise Adoption")
    ws3.append(["Distributor Name", "Total", "Installed", "Not Installed", "Uninstalled",
                 "Other", "Adoption Rate (%)"])
    if "Distributor Name" in df.columns:
        dist = df.groupby("Distributor Name").agg(
            total=("Account Sf ID", "count"),
            inst=("Status", lambda x: (x == "Installed").sum()),
            ni=("Status", lambda x: (x == "Not Installed").sum()),
            un=("Status", lambda x: (x == "Uninstalled").sum()),
            oth=("Status", lambda x: x.isin(["Tech Issue", "Not Working"]).sum()),
        ).reset_index()
        dist["rate"] = (dist["inst"] / dist["total"] * 100).round(1)
        for _, r in dist.sort_values("rate", ascending=False).iterrows():
            ws3.append([r["Distributor Name"], int(r["total"]), int(r["inst"]),
                        int(r["ni"]), int(r["un"]), int(r["oth"]), f"{r['rate']}%"])
    style_header(ws3)
    auto_width(ws3)

    # ── Sheet 4: Dealer Detail ──
    ws4 = wb.create_sheet("Dealer Detail")
    detail_cols = ["Account Sf ID", "Name of the Dealer", "State", "Zone",
                   "Distributor Name", "Account Owner As per SF", "Mobile No.",
                   "# orders (total)", "# orders (via. app.)", "Status"]
    avail = [c for c in detail_cols if c in df.columns]
    ws4.append(avail)
    status_order = {"Installed": 0, "Uninstalled": 1, "Tech Issue": 2, "Not Working": 3, "Not Installed": 4}
    sdf = df.copy()
    sdf["_s"] = sdf["Status"].map(status_order).fillna(5)
    sdf = sdf.sort_values(["_s", "Name of the Dealer"])
    for _, row in sdf.iterrows():
        ws4.append([row.get(c, "") if pd.notna(row.get(c)) else "" for c in avail])
    style_header(ws4)
    auto_width(ws4)

    # ── Sheet 5: App Usage ──
    ws5 = wb.create_sheet("App Usage (Orders)")
    ws5.append(["Account Sf ID", "Name of the Dealer", "Zone", "Distributor Name",
                "# Orders (Total)", "# Orders (via App)", "% Orders via App", "Status"])
    if "# orders (via. app.)" in df.columns:
        app_u = df[df["# orders (via. app.)"] > 0].sort_values("# orders (via. app.)", ascending=False)
        for _, r in app_u.iterrows():
            pct = f"{r['% orders via. app.']:.1f}%" if pd.notna(r.get("% orders via. app.")) else ""
            ws5.append([
                r.get("Account Sf ID", ""), r.get("Name of the Dealer", ""),
                r.get("Zone", ""), r.get("Distributor Name", ""),
                r.get("# orders (total)", 0), r.get("# orders (via. app.)", 0),
                pct, r.get("Status", ""),
            ])
    style_header(ws5)
    auto_width(ws5)

    # ── Sheet 6: At Risk ──
    ws6 = wb.create_sheet("At Risk - Action Needed")
    ws6.append(["Account Sf ID", "Name of the Dealer", "State", "Zone",
                "Distributor Name", "Account Owner As per SF", "Mobile No.",
                "Status", "Action Needed"])
    risk_map = {"Uninstalled": "Re-engage: app was uninstalled",
                "Tech Issue": "Resolve technical issue",
                "Not Working": "Fix: app not working"}
    risk = df[df["Status"].isin(risk_map.keys())].copy()
    risk["Action Needed"] = risk["Status"].map(risk_map)
    risk = risk.sort_values("Status")
    status_fills = {
        "Not Working": PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid"),
        "Tech Issue": PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid"),
        "Uninstalled": PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid"),
    }
    for _, r in risk.iterrows():
        ws6.append([
            r.get("Account Sf ID", ""), r.get("Name of the Dealer", ""),
            r.get("State", ""), r.get("Zone", ""),
            r.get("Distributor Name", ""), r.get("Account Owner As per SF", ""),
            r.get("Mobile No.", ""), r.get("Status", ""), r.get("Action Needed", ""),
        ])
    for ri in range(2, ws6.max_row + 1):
        sv = ws6.cell(row=ri, column=8).value
        if sv in status_fills:
            for ci in range(1, 10):
                ws6.cell(row=ri, column=ci).fill = status_fills[sv]
    style_header(ws6)
    auto_width(ws6)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════
# ── STREAMLIT UI ──
# ══════════════════════════════════════════════════════════════

st.title("📊 App Adoption Dashboard")

# ── Sidebar: File Upload ──
with st.sidebar:
    st.header("Upload Data")
    uploaded_file = st.file_uploader(
        "Upload your Excel file",
        type=["xlsx", "xls"],
        help="Upload the app adoption Excel file with dealer data",
    )

    if uploaded_file:
        st.success(f"Loaded: {uploaded_file.name}")

    st.markdown("---")
    st.markdown("**Expected columns:**")
    st.markdown("- Account Sf ID *(required)*\n- Name of the Dealer *(required)*\n- Status *(required)*\n- State, Zone, Distributor Name\n- Order data columns")

# ── Main Content ──
if uploaded_file is None:
    st.info("👈 Upload your Excel file in the sidebar to get started.")
    st.stop()

# Load data
raw_bytes = uploaded_file.getvalue()
df = load_and_clean(raw_bytes)

# Validate
is_valid, errors = validate_data(df)
if not is_valid:
    st.error("**Validation errors:**")
    for e in errors:
        st.write(f"- {e}")
    st.stop()

# ── Summary Metrics ──
total = len(df)
status_counts = df["Status"].value_counts()

installed = int(status_counts.get("Installed", 0))
not_installed = int(status_counts.get("Not Installed", 0))
uninstalled = int(status_counts.get("Uninstalled", 0))
tech_issue = int(status_counts.get("Tech Issue", 0))
not_working = int(status_counts.get("Not Working", 0))
adoption_rate = round(installed / total * 100, 1) if total > 0 else 0

st.header("Summary")

col1, col2, col3, col4, col5 = st.columns(5)
col1.metric("Total Dealers", total)
col2.metric("Installed", installed, f"{adoption_rate}%")
col3.metric("Not Installed", not_installed)
col4.metric("Uninstalled", uninstalled)
col5.metric("Tech Issue / Not Working", tech_issue + not_working)

st.markdown("---")

# ── Charts Row ──
chart_col1, chart_col2 = st.columns(2)

with chart_col1:
    st.subheader("Status Breakdown")
    status_df = pd.DataFrame({
        "Status": VALID_STATUSES,
        "Count": [int(status_counts.get(s, 0)) for s in VALID_STATUSES],
    })
    status_df = status_df[status_df["Count"] > 0]
    st.bar_chart(status_df.set_index("Status")["Count"], horizontal=True)

with chart_col2:
    st.subheader("Zone-wise Adoption")
    if "Zone" in df.columns:
        zone_df = df.groupby("Zone").agg(
            Total=("Account Sf ID", "count"),
            Installed=("Status", lambda x: (x == "Installed").sum()),
        ).reset_index()
        zone_df["Adoption %"] = (zone_df["Installed"] / zone_df["Total"] * 100).round(1)
        zone_df = zone_df.sort_values("Adoption %", ascending=False)

        st.dataframe(
            zone_df[["Zone", "Total", "Installed", "Adoption %"]],
            hide_index=True,
            use_container_width=True,
        )
    else:
        st.write("No Zone column found.")

st.markdown("---")

# ── Tabs for Details ──
tab1, tab2, tab3, tab4 = st.tabs([
    "📋 Distributor-wise", "👥 Dealer Detail", "📱 App Usage (Orders)", "⚠️ At Risk"
])

with tab1:
    st.subheader("Distributor-wise Adoption")
    if "Distributor Name" in df.columns:
        dist_df = df.groupby("Distributor Name").agg(
            Total=("Account Sf ID", "count"),
            Installed=("Status", lambda x: (x == "Installed").sum()),
            Uninstalled=("Status", lambda x: (x == "Uninstalled").sum()),
            Not_Installed=("Status", lambda x: (x == "Not Installed").sum()),
        ).reset_index()
        dist_df["Adoption %"] = (dist_df["Installed"] / dist_df["Total"] * 100).round(1)
        dist_df = dist_df.sort_values("Adoption %", ascending=False)
        dist_df.columns = ["Distributor Name", "Total", "Installed", "Uninstalled", "Not Installed", "Adoption %"]

        st.dataframe(dist_df, hide_index=True, use_container_width=True)
    else:
        st.write("No Distributor Name column found.")

with tab2:
    st.subheader("All Dealers")

    # Filters
    fcol1, fcol2, fcol3 = st.columns(3)
    with fcol1:
        status_filter = st.multiselect("Filter by Status", VALID_STATUSES, default=VALID_STATUSES)
    with fcol2:
        zones = sorted(df["Zone"].dropna().unique()) if "Zone" in df.columns else []
        zone_filter = st.multiselect("Filter by Zone", zones, default=zones)
    with fcol3:
        search = st.text_input("Search dealer name")

    filtered = df[df["Status"].isin(status_filter)]
    if "Zone" in df.columns and zone_filter:
        filtered = filtered[filtered["Zone"].isin(zone_filter)]
    if search:
        filtered = filtered[
            filtered["Name of the Dealer"].fillna("").str.contains(search, case=False, na=False)
        ]

    display_cols = ["Account Sf ID", "Name of the Dealer", "State", "Zone",
                    "Distributor Name", "Account Owner As per SF", "Mobile No.",
                    "# orders (total)", "# orders (via. app.)", "Status"]
    avail_cols = [c for c in display_cols if c in filtered.columns]

    st.write(f"Showing **{len(filtered)}** of {total} dealers")
    st.dataframe(filtered[avail_cols], hide_index=True, use_container_width=True, height=500)

with tab3:
    st.subheader("Dealers Placing Orders via App")
    if "# orders (via. app.)" in df.columns:
        app_users = df[df["# orders (via. app.)"] > 0].sort_values(
            "# orders (via. app.)", ascending=False
        )
        if len(app_users) == 0:
            st.info("No dealers have placed orders via the app yet.")
        else:
            st.write(f"**{len(app_users)}** dealers have placed orders through the app")
            order_cols = ["Account Sf ID", "Name of the Dealer", "Zone", "Distributor Name",
                          "# orders (total)", "# orders (via. app.)", "% orders via. app.", "Status"]
            avail_order_cols = [c for c in order_cols if c in app_users.columns]
            st.dataframe(app_users[avail_order_cols], hide_index=True, use_container_width=True)
    else:
        st.write("No order data columns found.")

with tab4:
    st.subheader("Action Needed")
    st.write("Dealers who have uninstalled, have tech issues, or app not working.")

    risk_map = {
        "Not Working": "🔴 Fix: app not working",
        "Tech Issue": "🟠 Resolve technical issue",
        "Uninstalled": "🟡 Re-engage: app was uninstalled",
    }

    risk_df = df[df["Status"].isin(risk_map.keys())].copy()
    risk_df["Action"] = risk_df["Status"].map(risk_map)

    rcol1, rcol2, rcol3 = st.columns(3)
    rcol1.metric("🔴 Not Working", not_working)
    rcol2.metric("🟠 Tech Issue", tech_issue)
    rcol3.metric("🟡 Uninstalled", uninstalled)

    if len(risk_df) > 0:
        risk_cols = ["Account Sf ID", "Name of the Dealer", "State", "Zone",
                     "Distributor Name", "Account Owner As per SF", "Mobile No.",
                     "Status", "Action"]
        avail_risk = [c for c in risk_cols if c in risk_df.columns]
        st.dataframe(
            risk_df[avail_risk].sort_values("Status"),
            hide_index=True, use_container_width=True, height=400,
        )
    else:
        st.success("No at-risk dealers!")

# ── Download Report ──
st.markdown("---")
st.header("Download Report")
st.write("Generate a formatted Excel report with all the data above.")

report_bytes = generate_excel_report(df)
st.download_button(
    label="📥 Download Excel Report",
    data=report_bytes,
    file_name=f"adoption_report_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary",
)
