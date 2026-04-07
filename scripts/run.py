#!/usr/bin/env python3
"""
CLI entry point for the App Adoption Tracking System.

Usage:
    python scripts/run.py ingest <file>
    python scripts/run.py report
    python scripts/run.py validate <file>
    python scripts/run.py create-template
"""

import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from scripts.config import (
    get_excel_columns, SAMPLE_DATA, TEMPLATE_DIR, VALID_STATUSES,
)
from scripts.validate import validate_upload
from scripts.ingest import ingest_file
from scripts.reports import generate_reports


def cmd_create_template(_args):
    """Create an Excel template matching the expected data format."""
    os.makedirs(TEMPLATE_DIR, exist_ok=True)

    template_path = os.path.join(TEMPLATE_DIR, "app_adoption_template.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    columns = get_excel_columns()

    # Header row
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Sample data rows
    for row_idx, sample in enumerate(SAMPLE_DATA, 2):
        for col_idx, col_name in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=sample.get(col_name, ""))

    # Status dropdown validation
    status_col = columns.index("Status") + 1
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(VALID_STATUSES)}"',
        allow_blank=False,
    )
    dv.error = f"Please enter one of: {', '.join(VALID_STATUSES)}"
    ws.add_data_validation(dv)
    for row_idx in range(2, 1002):
        dv.add(ws.cell(row=row_idx, column=status_col))

    # Auto-fit column widths
    for col_idx, col_name in enumerate(columns, 1):
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"].width = max(len(col_name) + 4, 15)

    ws.freeze_panes = "A2"

    # Instructions sheet
    ws_instr = wb.create_sheet("Instructions")
    instructions = [
        "App Adoption Data Template",
        "",
        "How to use:",
        "1. Go to the 'Data' sheet",
        "2. Replace the sample rows with your actual data",
        "3. Do NOT rename or reorder the columns",
        "4. Save the file",
        "5. Run: python scripts/run.py ingest <this_file.xlsx>",
        "",
        "Required columns:",
        "  - Account Sf ID (unique identifier for each dealer)",
        "  - Name of the Dealer",
        "  - Status (Installed / Not Installed / Uninstalled / Tech Issue / Not Working)",
        "",
        "All other columns are optional but recommended.",
        "",
        "For order columns, use '-' or leave blank if no data.",
    ]
    for row_idx, line in enumerate(instructions, 1):
        cell = ws_instr.cell(row=row_idx, column=1, value=line)
        if row_idx == 1:
            cell.font = Font(bold=True, size=14)
    ws_instr.column_dimensions["A"].width = 70

    wb.save(template_path)
    print(f"Template created: {template_path}")


def cmd_validate(args):
    """Validate an uploaded Excel file."""
    if not os.path.exists(args.file):
        print(f"ERROR: File not found: {args.file}")
        sys.exit(1)

    is_valid, messages = validate_upload(args.file)

    print(f"\nValidation: {os.path.basename(args.file)}")
    print("-" * 50)
    for msg in messages:
        print(f"  {msg}")
    print("-" * 50)

    if is_valid:
        print("RESULT: PASSED")
        print(f"Next: python scripts/run.py ingest \"{args.file}\"")
    else:
        print("RESULT: FAILED - fix errors above and retry.")
        sys.exit(1)


def cmd_ingest(args):
    """Ingest an uploaded Excel file."""
    if not os.path.exists(args.file):
        print(f"ERROR: File not found: {args.file}")
        sys.exit(1)

    print(f"Ingesting: {os.path.basename(args.file)}")
    print("-" * 50)

    try:
        rows_affected, messages = ingest_file(args.file)
        for msg in messages:
            print(f"  {msg}")
        print("-" * 50)
        print(f"SUCCESS: {rows_affected} row(s) processed.")
        print("Next: python scripts/run.py report")
    except ValueError as e:
        print(f"\n{e}")
        sys.exit(1)


def cmd_report(_args):
    """Generate adoption reports."""
    print("Generating adoption report...")
    print("-" * 50)

    try:
        output_path = generate_reports()
        print(f"  Report saved to: {output_path}")
        print("-" * 50)
        print("SUCCESS: Open the Excel file to view the report.")
    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)


def main():
    parser = argparse.ArgumentParser(
        description="App Adoption Tracking System",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python scripts/run.py ingest "data/App adoption till 6th Apr.xlsx"
  python scripts/run.py report
  python scripts/run.py validate "data/myfile.xlsx"
  python scripts/run.py create-template
        """,
    )
    subparsers = parser.add_subparsers(dest="command", help="Command to run")

    sub = subparsers.add_parser("create-template", help="Create Excel template")
    sub.set_defaults(func=cmd_create_template)

    sub = subparsers.add_parser("validate", help="Validate an Excel file")
    sub.add_argument("file", help="Path to the Excel file")
    sub.set_defaults(func=cmd_validate)

    sub = subparsers.add_parser("ingest", help="Import an Excel file")
    sub.add_argument("file", help="Path to the Excel file")
    sub.set_defaults(func=cmd_ingest)

    sub = subparsers.add_parser("report", help="Generate adoption report")
    sub.set_defaults(func=cmd_report)

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        sys.exit(1)

    args.func(args)


if __name__ == "__main__":
    main()
